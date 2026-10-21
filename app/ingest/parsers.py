import os
import re

HEADING = re.compile(r"^\s*(#{1,4}\s+.+|\d+\.(\d+)*\s+\S.*|§\s*\d+.*|Section\s+\d+\b.*)$", re.I)


class Segment:
    def __init__(self, text: str, page: int | None, heading: str | None):
        self.text = text.strip()
        self.page = page
        self.heading = heading


def _looks_like_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 90:
        return False
    return bool(HEADING.match(line))


def _segments_from_lines(lines: list[str], page: int | None) -> list[Segment]:
    segments, heading, buf = [], None, []
    for line in lines:
        if _looks_like_heading(line):
            if buf:
                segments.append(Segment("\n".join(buf), page, heading))
                buf = []
            heading = line.strip().lstrip("#").strip()
        else:
            buf.append(line)
    if buf:
        segments.append(Segment("\n".join(buf), page, heading))
    return [s for s in segments if s.text]


def parse_pdf(path: str) -> tuple[list[Segment], int]:
    import pymupdf

    doc = pymupdf.open(path)
    text_chars = sum(len(p.get_text()) for p in doc)
    if text_chars < 100:
        return _ocr_pdf(doc), doc.page_count
    segments = []
    for i, page in enumerate(doc, start=1):
        segments += _segments_from_lines(page.get_text().splitlines(), i)
    return segments, doc.page_count


def _ocr_pdf(doc) -> list[Segment]:
    import pytesseract
    from PIL import Image

    segments = []
    for i, page in enumerate(doc, start=1):
        pix = page.get_pixmap(dpi=300)
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        text = pytesseract.image_to_string(img)
        segments += _segments_from_lines(text.splitlines(), i)
    return segments


def parse_docx(path: str) -> tuple[list[Segment], int]:
    import docx

    document = docx.Document(path)
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    segments = _segments_from_lines(lines, None)
    for table in document.tables:
        rows = [" | ".join(c.text for c in row.cells) for row in table.rows]
        segments.append(Segment("\n".join(rows), None, "Table"))
    return segments, 0


def parse_html(path: str) -> tuple[list[Segment], int]:
    from bs4 import BeautifulSoup

    with open(path, encoding="utf-8") as f:
        soup = BeautifulSoup(f, "lxml")
    lines = soup.get_text("\n").splitlines()
    return _segments_from_lines([l for l in lines if l.strip()], None), 0


def parse_md(path: str) -> tuple[list[Segment], int]:
    with open(path, encoding="utf-8") as f:
        return _segments_from_lines(f.read().splitlines(), None), 0


def parse(path: str) -> tuple[list[Segment], int]:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return parse_pdf(path)
    if ext == ".docx":
        return parse_docx(path)
    if ext in (".html", ".htm"):
        return parse_html(path)
    if ext in (".md", ".markdown"):
        return parse_md(path)
    raise ValueError(f"unsupported format: {ext}")
